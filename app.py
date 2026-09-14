import os
import datetime
from flask import Flask, render_template, request, redirect, url_for, flash
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename
from flask import Flask, render_template, request, redirect, url_for, flash, send_file

app = Flask(__name__)
app.secret_key = "esports_club_key"

# Paths
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, 'uploads')
EXCEL_FILE = os.path.join(BASE_DIR, 'Tournament_Registrations.xlsx')
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'pdf'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def init_excel():
    try:
        if not os.path.exists(EXCEL_FILE):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Registrations"
            ws.views.sheetView[0].showGridLines = True

            headers = [
                "Timestamp", "Game", "Team / Clan Name", "Team Leader Name", 
                "Admission No", "Branch & Section", "In-Game ID / UID", 
                "WhatsApp Number", "UPI Transaction ID (UTR)", "Payment Screenshot File"
            ]

            header_fill = PatternFill(start_color="1A2536", end_color="1A2536", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="00FFCC")
            
            ws.append(headers)
            ws.row_dimensions[1].height = 26

            thin_border = Border(
                left=Side(style='thin', color='D0D7DE'),
                right=Side(style='thin', color='D0D7DE'),
                top=Side(style='thin', color='D0D7DE'),
                bottom=Side(style='thin', color='D0D7DE')
            )

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
                ws.column_dimensions[get_column_letter(col_idx)].width = 22

            wb.save(EXCEL_FILE)
            wb.close()
    except Exception as e:
        print(f"Error initializing Excel: {e}")

# Initialize on startup
init_excel()

# Route 1: Home Page
@app.route('/')
def home():
    return render_template('index.html')

# Route 2: Dedicated registration pages for BGMI and Free Fire
@app.route('/register/<game_type>')
def register_page(game_type):
    game_type = game_type.lower()
    if game_type == 'bgmi':
        game_title = "BGMI (Battlegrounds Mobile India)"
        bg_class = "bgmi-bg"
    elif game_type in ['freefire', 'free-fire']:
        game_title = "Free Fire MAX"
        bg_class = "ff-bg"
    else:
        return redirect(url_for('home'))
        
    return render_template('register.html', game_title=game_title, bg_class=bg_class, game_slug=game_type)

# Route 3: Form submission handler (Auto-saves to Excel)
@app.route('/submit-registration', methods=['POST'])
def submit_registration():
    game = request.form.get('game')
    team_name = request.form.get('team_name')
    full_name = request.form.get('name')
    admission_no = request.form.get('admission_no')
    branch = request.form.get('branch')
    in_game_id = request.form.get('in_game_id')
    contact = request.form.get('contact')
    utr_id = request.form.get('utr_id')
    game_slug = request.form.get('game_slug', 'bgmi')

    # Save payment screenshot
    file = request.files.get('payment_screenshot')
    if not file or file.filename == '' or not allowed_file(file.filename):
        flash("Invalid file format. Please upload JPG, PNG, or PDF.")
        return redirect(url_for('register_page', game_type=game_slug))

    timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = secure_filename(file.filename)
    filename = f"{admission_no}_{timestamp_str}_{safe_name}"
    file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(file_path)

    timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Append directly to your Excel file
    try:
        init_excel()
        wb = openpyxl.load_workbook(EXCEL_FILE)
        ws = wb["Registrations"]

        row_data = [
            timestamp, game, team_name, full_name, admission_no,
            branch, in_game_id, contact, utr_id, filename
        ]
        ws.append(row_data)

        new_row_idx = ws.max_row
        thin_border = Border(
            left=Side(style='thin', color='D0D7DE'),
            right=Side(style='thin', color='D0D7DE'),
            top=Side(style='thin', color='D0D7DE'),
            bottom=Side(style='thin', color='D0D7DE')
        )

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=new_row_idx, column=col_idx)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if col_idx in [1, 2, 5, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        wb.save(EXCEL_FILE)
        wb.close()
        flash(f"Slot registered successfully for {game}!")
    except Exception as e:
        print(f"CRITICAL Error saving to Excel: {e}")
        flash("Server error while saving details. Please contact support.")

    return redirect(url_for('register_page', game_type=game_slug))

# Secret Admin Route to Download Live Excel Registrations from Render
@app.route('/admin/download-excel')
def download_excel():
    if os.path.exists(EXCEL_FILE):
        return send_file(
            EXCEL_FILE, 
            as_attachment=True, 
            download_name="Tournament_Registrations.xlsx"
        )
    return "No registrations found yet or file was reset by Render.", 404

if __name__ == '__main__':
    app.run(debug=True, port=5000)
    
    